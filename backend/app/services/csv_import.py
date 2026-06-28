import csv
import re
from datetime import datetime
from io import BytesIO, StringIO
from typing import Any

from openpyxl import load_workbook


REQUIRED_FIELDS = {
    "transaction_reference",
    "amount",
    "status",
    "payment_date",
    "customer_identifier",
    "provider",
}

GENERAL_GUIDANCE = (
    "Please upload a bank or payment export in Excel or CSV format. "
    "For best results, export your statement directly from your bank, OPay, "
    "Paystack, or Flutterwave dashboard."
)

SUCCESS_GUIDANCE = "You can now run reconciliation, view dashboard, or export your report."

BANK_STATEMENT_HEADER_KEYWORDS = {
    "trans date",
    "transaction date",
    "date",
    "value date",
    "description",
    "narration",
    "details",
    "debit",
    "debit ₦",
    "debit(₦)",
    "credit",
    "credit ₦",
    "credit(₦)",
    "amount",
    "balance",
    "balance after",
    "balance after ₦",
    "channel",
    "transaction reference",
    "reference",
    "ref",
    "chq no",
}

COLUMN_MAPPINGS = {
    "payment_date": {
        "trans date",
        "transaction date",
        "payment date",
        "date",
        "value date",
        "paid at",
        "created at",
        "time",
    },
    "customer_identifier": {
        "description",
        "narration",
        "details",
        "customer",
        "customer name",
        "customer identifier",
        "email",
        "customer email",
        "phone",
        "customer phone",
        "payer",
        "account name",
    },
    "transaction_reference": {
        "transaction reference",
        "reference",
        "ref",
        "transaction ref",
        "transaction id",
        "payment reference",
        "provider reference",
        "chq no",
    },
    "amount": {
        "amount",
        "amount paid",
        "paid amount",
        "transaction amount",
        "value",
    },
    "debit": {
        "debit",
        "debit ₦",
        "debit(₦)",
    },
    "credit": {
        "credit",
        "credit ₦",
        "credit(₦)",
    },
    "status": {
        "status",
        "payment status",
        "transaction status",
        "result",
    },
    "provider": {
        "provider",
        "payment provider",
        "gateway",
        "channel",
        "source",
    },
}


def normalize_column_name(column_name: Any) -> str:
    cleaned = str(column_name or "").strip().lower()
    cleaned = cleaned.replace("_", " ").replace("-", " ")
    cleaned = cleaned.replace("\u20a6", "₦")
    cleaned = re.sub(r"\s+", " ", cleaned)
    return cleaned


def friendly_file_type(extension: str) -> str:
    if extension == "xlsx":
        return "Excel"
    if extension == "csv":
        return "CSV"
    if extension == "pdf":
        return "PDF"
    if extension in {"png", "jpg", "jpeg"}:
        return "Image"
    return extension.upper() if extension else "Unknown"


def get_file_extension(filename: str) -> str:
    if not filename or "." not in filename:
        return ""
    return filename.rsplit(".", 1)[-1].lower()


def clean_text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def clean_money(value: Any) -> float | None:
    text = clean_text(value)
    if not text or text in {"-", "--"}:
        return None

    cleaned = (
        text.replace("\u20a6", "")
        .replace("₦", "")
        .replace(",", "")
        .replace(" ", "")
    )

    if cleaned in {"", "-", "--"}:
        return None

    try:
        return float(cleaned)
    except ValueError:
        return None


def normalize_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()

    text = clean_text(value)
    if not text:
        return ""

    formats = [
        "%Y-%m-%d",
        "%Y-%m-%d %H:%M:%S",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%m/%d/%Y",
        "%d %b %Y",
        "%d %B %Y",
    ]

    for date_format in formats:
        try:
            return datetime.strptime(text, date_format).date().isoformat()
        except ValueError:
            continue

    return text


def base_response(
    *,
    message: str,
    imported: int = 0,
    rejected: int = 0,
    detected_file_type: str = "unknown",
    detected_statement_type: str = "unsupported",
    detected_provider: str = "Bank Statement",
    detected_columns: list[str] | None = None,
    mapped_columns: dict[str, str] | None = None,
    missing_required_fields: list[str] | None = None,
    rejected_rows: list[dict[str, Any]] | None = None,
    valid_rows: list[dict[str, Any]] | None = None,
    user_guidance: str = GENERAL_GUIDANCE,
):
    rows = rejected_rows or []

    return {
        "message": message,
        "imported": imported,
        "rejected": rejected,
        "total_rejected_rows": len(rows) if rejected else rejected,
        "detected_file_type": detected_file_type,
        "detected_statement_type": detected_statement_type,
        "detected_provider": detected_provider,
        "detected_columns": detected_columns or [],
        "mapped_columns": mapped_columns or {},
        "missing_required_fields": missing_required_fields or [],
        "rejected_rows": rows,
        "valid_rows": valid_rows or [],
        "user_guidance": user_guidance,
    }


def read_csv_rows(content: bytes):
    decoded_content = content.decode("utf-8-sig")
    reader = csv.DictReader(StringIO(decoded_content))
    detected_columns = reader.fieldnames or []
    return detected_columns, list(reader), "transaction_table"


def row_has_transaction_headers(row_values) -> bool:
    normalized_values = {
        normalize_column_name(value)
        for value in row_values
        if clean_text(value)
    }

    matches = sum(
        1
        for known_header in BANK_STATEMENT_HEADER_KEYWORDS
        if known_header in normalized_values
    )

    return matches >= 3


def find_xlsx_header_row(rows) -> int:
    # Search first 30 rows because real bank statements often contain metadata above headers.
    for index, row_values in enumerate(rows[:30]):
        if row_has_transaction_headers(row_values):
            return index
    return 0


def read_xlsx_rows(content: bytes):
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))

    if not rows:
        return [], [], "unsupported"

    header_row_index = find_xlsx_header_row(rows)
    detected_columns = [
        clean_text(value)
        for value in rows[header_row_index]
    ]

    data_rows = []

    for row_values in rows[header_row_index + 1:]:
        if not any(clean_text(value) for value in row_values):
            continue

        row = {}
        for index, column in enumerate(detected_columns):
            row[column] = row_values[index] if index < len(row_values) else None
        data_rows.append(row)

    statement_type = "bank_statement" if header_row_index > 0 else "transaction_table"
    return detected_columns, data_rows, statement_type


def map_columns(detected_columns):
    mapped_columns = {}

    for original_column in detected_columns:
        normalized = normalize_column_name(original_column)

        for internal_field, variations in COLUMN_MAPPINGS.items():
            if normalized in variations:
                mapped_columns[original_column] = internal_field
                break

    return mapped_columns


def get_first_value(raw_row: dict[str, Any], mapped_columns: dict[str, str], target_field: str) -> Any:
    for original_column, internal_field in mapped_columns.items():
        if internal_field == target_field:
            value = raw_row.get(original_column)
            if clean_text(value):
                return value
    return ""


def normalize_transaction_row(
    *,
    raw_row: dict[str, Any],
    mapped_columns: dict[str, str],
    row_number: int,
    provider_fallback: str | None,
    detected_statement_type: str,
):
    normalized_row = {}
    row_errors = []

    for original_column, internal_field in mapped_columns.items():
        if internal_field in {"debit", "credit"}:
            continue

        value = raw_row.get(original_column)
        normalized_row[internal_field] = clean_text(value)

    payment_date = normalize_date(normalized_row.get("payment_date"))
    normalized_row["payment_date"] = payment_date

    amount_value = clean_money(normalized_row.get("amount"))

    if amount_value is None and detected_statement_type == "bank_statement":
        credit_value = clean_money(get_first_value(raw_row, mapped_columns, "credit"))
        debit_value = clean_money(get_first_value(raw_row, mapped_columns, "debit"))

        if credit_value is not None:
            amount_value = credit_value
            normalized_row["transaction_direction"] = "credit"
        elif debit_value is not None:
            amount_value = debit_value
            normalized_row["transaction_direction"] = "debit"

    if amount_value is None:
        row_errors.append("Missing amount")
    else:
        normalized_row["amount"] = amount_value

    description = clean_text(normalized_row.get("customer_identifier"))

    if detected_statement_type == "bank_statement":
        normalized_row["status"] = normalized_row.get("status") or "success"
        normalized_row["provider"] = normalized_row.get("provider") or provider_fallback or "Bank Statement"

        if not normalized_row.get("customer_identifier") and description:
            normalized_row["customer_identifier"] = description

        if not normalized_row.get("transaction_reference") and payment_date and amount_value is not None and description:
            safe_date = payment_date.replace("-", "")
            safe_amount = str(amount_value).replace(".", "")
            normalized_row["transaction_reference"] = f"BANK-{row_number}-{safe_date}-{safe_amount}"

    elif provider_fallback and not normalized_row.get("provider"):
        normalized_row["provider"] = provider_fallback

    for required_field in REQUIRED_FIELDS:
        if required_field == "amount" and "Missing amount" in row_errors:
            continue

        if not normalized_row.get(required_field):
            friendly_name = required_field.replace("_", " ")
            row_errors.append(f"Missing {friendly_name}")

    return normalized_row, row_errors


def parse_smart_transaction_file(
    filename: str,
    content: bytes,
    provider_fallback: str | None = None,
):
    extension = get_file_extension(filename)

    if extension == "xls":
        return base_response(
            message="We could not read this old Excel file clearly.",
            detected_file_type="xls",
            user_guidance="Please save or export the file as .xlsx, then upload it again.",
        )

    if extension == "pdf":
        return base_response(
            message="We could not read this PDF clearly.",
            detected_file_type="pdf",
        )

    if extension in {"png", "jpg", "jpeg"}:
        return base_response(
            message="We could not read this screenshot clearly.",
            detected_file_type=extension,
        )

    try:
        if extension == "csv":
            detected_columns, raw_rows, detected_statement_type = read_csv_rows(content)
        elif extension == "xlsx":
            detected_columns, raw_rows, detected_statement_type = read_xlsx_rows(content)
        else:
            return base_response(
                message="We could not read this file clearly.",
                detected_file_type=extension or "unknown",
            )
    except Exception:
        return base_response(
            message="We could not read this file clearly.",
            detected_file_type=extension or "unknown",
        )

    mapped_columns = map_columns(detected_columns)
    mapped_fields = set(mapped_columns.values())

    if detected_statement_type == "bank_statement":
        mapped_fields.add("status")
        mapped_fields.add("provider")

        if "debit" in mapped_fields or "credit" in mapped_fields:
            mapped_fields.add("amount")

    if provider_fallback and "provider" not in mapped_fields:
        mapped_fields.add("provider")

    missing_required_fields = sorted(
        field for field in REQUIRED_FIELDS if field not in mapped_fields
    )

    valid_rows = []
    rejected_rows = []

    if not detected_columns or len(mapped_fields) < 3:
        return base_response(
            message="We could not read this file clearly.",
            detected_file_type=extension,
            detected_statement_type="unsupported",
            detected_columns=detected_columns,
            mapped_columns=mapped_columns,
            missing_required_fields=missing_required_fields,
        )

    for index, raw_row in enumerate(raw_rows, start=2):
        normalized_row, row_errors = normalize_transaction_row(
            raw_row=raw_row,
            mapped_columns=mapped_columns,
            row_number=index,
            provider_fallback=provider_fallback,
            detected_statement_type=detected_statement_type,
        )

        if row_errors:
            rejected_rows.append({
                "row": index,
                "reason": "; ".join(row_errors),
            })
            continue

        valid_rows.append(normalized_row)

    statement_provider = provider_fallback or (
        "Bank Statement" if detected_statement_type == "bank_statement" else "Smart Import"
    )

    return base_response(
        message="File processed successfully.",
        imported=len(valid_rows),
        rejected=len(rejected_rows),
        detected_file_type=extension,
        detected_statement_type=detected_statement_type,
        detected_provider=statement_provider,
        detected_columns=detected_columns,
        mapped_columns=mapped_columns,
        missing_required_fields=missing_required_fields,
        rejected_rows=rejected_rows,
        valid_rows=valid_rows,
        user_guidance=SUCCESS_GUIDANCE,
    )


def parse_transaction_csv(file_content: str):
    result = parse_smart_transaction_file(
        "transactions.csv",
        file_content.encode("utf-8"),
    )

    if result["missing_required_fields"]:
        return result["valid_rows"], [
            f"Missing columns: {', '.join(result['missing_required_fields'])}",
        ]

    errors = [
        f"Row {row['row']}: {row['reason']}"
        for row in result["rejected_rows"]
    ]

    return result["valid_rows"], errors
